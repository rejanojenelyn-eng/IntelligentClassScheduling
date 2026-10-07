"""
Faculty side → Weekly Schedule export (PDF / XLSX / CSV / DOCX).


Same document design as the Teaching Assignment export (teaching_assignment_export):
centred PUP logo, bold black title + term line, the maroon "LOPEZ, QUEZON
CAMPUS" banner, the EMP NO / EMP NAME / EMP STATUS × COLLEGE / DEPT CODE /
DEPARTMENT block and black-grid tables — adapted to a week:


  1. WEEKLY TIMETABLE — 30-minute rows (7:30 AM – 9:30 PM) × MON–SUN, each
     day headed with its date; a class fills the rows it occupies. Days of the
     week outside the semester are marked "No classes".
  2. CLASS LIST — every meeting that week: Day, Date, Time, Subject, Year &
     Section, Room, Hours, Remarks (approved make-up classes are flagged).
  3. HOURS PER DAY — total teaching hours for each day of that week.


`d` comes from app.faculty_weekly_schedule_export: the Teaching Assignment
info fields plus week (7 dates), week_in_semester, sem_start/sem_end and
entries (recurring classes that fall inside the semester that week, plus that
week's approved make-ups, kind='makeup').
"""
import csv
import io
import re


import teaching_assignment_export as ta


GRID_START, GRID_END, SLOT = 7 * 60 + 30, 21 * 60 + 30, 30          # 7:30 AM – 9:30 PM
SLOTS = list(range(GRID_START, GRID_END, SLOT))                       # 28 half-hour rows
LIST_HEADERS = ['DAY', 'DATE', 'TIME', 'SUBJECT CODE', 'SUBJECT DESCRIPTION',
                'YEAR & SECTION', 'ROOM', 'HOURS', 'REMARKS']
LIST_CENTER = {0, 1, 2, 3, 5, 6, 7, 8}
MAKEUP_FILL, CLASS_FILL, OFF_FILL = 'FCE8E8', 'F2F2F2', 'D9D9D9'




# ─────────────────────────────────────────────────────────────
#  Shared model
# ─────────────────────────────────────────────────────────────
def _mins(t):
    m = re.match(r'\s*(\d{1,2}):(\d{2})\s*([AP]M)', t or '', re.I)
    if not m:
        return None
    h, mm, ap = int(m.group(1)), int(m.group(2)), m.group(3).upper()
    if ap == 'PM' and h != 12:
        h += 12
    if ap == 'AM' and h == 12:
        h = 0
    return h * 60 + mm




def _span(e):
    parts = str(e.get('time_range') or '').split(' - ')
    if len(parts) != 2:
        return None, None
    return _mins(parts[0]), _mins(parts[1])




def _slot_label(m):
    h, mm = divmod(m, 60)
    ap = 'AM' if h < 12 else 'PM'
    h12 = h - 12 if h > 12 else (12 if h == 0 else h)
    return f"{h12}:{mm:02d} {ap}"




def _fmt_date(d, fmt='%b %d'):
    return d.strftime(fmt).replace(' 0', ' ')




def _week_label(d):
    w = d.get('week') or []
    if not w:
        return ''
    a, b = w[0], w[-1]
    left = _fmt_date(a, '%b %d') + (f", {a.year}" if a.year != b.year else '')
    return f"WEEK OF {left.upper()} – {_fmt_date(b, '%b %d, %Y').upper()}"




def _title_lines(d):
    base = ta._title_lines(d)                      # ['FACULTY ASSIGNMENT', 'FIRST SEMESTER, SY …', (LOCAL)]
    lines = ['WEEKLY SCHEDULE', base[1], _week_label(d)]
    if d.get('source') == 'local':
        lines.append('LOCAL SCHEDULE')
    if d.get('program_filter'):
        lines.append(f"PROGRAM: {d['program_filter']}")
    return lines




def _entries(d):
    out = []
    for e in d.get('entries') or []:
        s, t = _span(e)
        out.append({**e, 'start': s, 'end': t})
    out.sort(key=lambda e: (e['date'], e['start'] or 0))
    return out




def _day_headers(d):
    return [f"{h}\n{_fmt_date(day)}" for h, day in zip(ta.DAY_HEADERS, d.get('week') or [])]




def _cell_text(e):
    bits = [e.get('subjectcode') or '', e.get('year_section') or '', e.get('room') or '']
    if e.get('kind') == 'makeup':
        bits.append('MAKE-UP')
    return '\n'.join(b for b in bits if b)




def grid_model(d):
    """{day_index: [(start_row, end_row_exclusive, text, kind)]} for the timetable."""
    raw = {i: [] for i in range(7)}          # [r0, r1, [entries]]
    week = d.get('week') or []
    for e in _entries(d):
        if e['start'] is None or e['end'] is None or e['date'] not in week:
            continue
        i = week.index(e['date'])
        r0 = max(0, (e['start'] - GRID_START) // SLOT)
        r1 = min(len(SLOTS), -(-(e['end'] - GRID_START) // SLOT))
        if r1 <= r0:
            continue
        clash = next((b for b in raw[i] if r0 < b[1] and b[0] < r1), None)
        if clash:   # overlapping meetings share one block rather than breaking the grid
            clash[0], clash[1] = min(clash[0], r0), max(clash[1], r1)
            clash[2].append(e)
        else:
            raw[i].append([r0, r1, [e]])
    blocks = {}
    for i, lst in raw.items():
        blocks[i] = []
        for r0, r1, es in lst:
            if len(es) == 1:
                text = _cell_text(es[0])
            else:   # one compact line per class so the block never overflows
                text = '\n'.join(f"{ta.compact_time_range(x.get('time_range'))} {x.get('subjectcode') or ''} "
                                 f"{x.get('year_section') or ''}"
                                 + (' (MAKE-UP)' if x.get('kind') == 'makeup' else '') for x in es)
            kind = 'makeup' if any(x.get('kind') == 'makeup' for x in es) else 'class'
            blocks[i].append((r0, r1, text, kind))
    return blocks




def list_rows(d):
    rows = []
    for e in _entries(d):
        rows.append([e['date'].strftime('%A'), _fmt_date(e['date'], '%m/%d/%Y').replace('/0', '/').lstrip('0'),
                     ta.compact_time_range(e.get('time_range')), e.get('subjectcode') or '',
                     e.get('subjectname') or '', e.get('year_section') or '', e.get('room') or '',
                     ta._num(e.get('hrs')),
                     'Make-up class' if e.get('kind') == 'makeup' else ''])
    return rows




def hours_row(d):
    week, ins = d.get('week') or [], d.get('week_in_semester') or [True] * 7
    per = {day: 0.0 for day in week}
    for e in _entries(d):
        if e['date'] in per:
            per[e['date']] += float(e.get('hrs') or 0)
    vals = [('—' if not ok else (ta._num(round(per[day], 2)) if per[day] else '')) for day, ok in zip(week, ins)]
    total = ta._num(round(sum(per.values()), 2))
    return vals, total




def notes(d):
    ins = d.get('week_in_semester') or [True] * 7
    s, e = d.get('sem_start'), d.get('sem_end')
    rng = f" ({_fmt_date(s, '%b %d, %Y')} – {_fmt_date(e, '%b %d, %Y')})" if s and e else ''
    out = []
    if not any(ins):
        out.append(f"This week is outside the {d.get('sem_label') or 'semester'}{rng}, so there are no classes.")
    elif not all(ins):
        out.append(f"Days marked \"No classes\" fall outside the {d.get('sem_label') or 'semester'}{rng}.")
    if any(x.get('kind') == 'makeup' for x in d.get('entries') or []):
        out.append('MAKE-UP: an approved make-up class held on that date only (not part of the regular weekly schedule).')
    if not (d.get('entries') or []) and any(ins):
        out.append('No classes are scheduled this week.')
    return out




# ─────────────────────────────────────────────────────────────
#  CSV
# ─────────────────────────────────────────────────────────────
def gen_csv(d, logo_path=None):
    out = io.StringIO()
    w = csv.writer(out)
    for line in _title_lines(d):
        w.writerow([line])
    w.writerow([ta.CAMPUS])
    w.writerow([])
    for row in ta._info_rows(d):
        w.writerow(list(row))
    w.writerow([])
    w.writerow(['CLASS LIST'])
    w.writerow(LIST_HEADERS)
    for r in list_rows(d):
        w.writerow(r)
    w.writerow([])
    w.writerow(['TEACHING HOURS PER DAY'])
    w.writerow([''] + [h.replace('\n', ' ') for h in _day_headers(d)] + ['TOTAL'])
    vals, total = hours_row(d)
    w.writerow(['HOURS'] + vals + [total])
    for n in notes(d):
        w.writerow([])
        w.writerow([n])
    return out.getvalue().encode('utf-8-sig')




# ─────────────────────────────────────────────────────────────
#  XLSX
# ─────────────────────────────────────────────────────────────
def gen_xlsx(d, logo_path=None):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter


    wb = Workbook()
    ws = wb.active
    ws.title = 'Weekly Schedule'
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = 'landscape'
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0


    fill = lambda hx: PatternFill('solid', fgColor=hx)
    MAROON, SHADE = fill(ta.MAROON_HEX), fill('F2F2F2')
    thin = Side(style='thin', color='000000')
    brd = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    left = Alignment(horizontal='left', vertical='center', wrap_text=True)
    NC = 9                                          # time + 7 days + list overflow column
    widths = [11, 16, 16, 16, 16, 16, 16, 16, 18]
    for ci, wdt in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(ci)].width = wdt


    def merged(r, c1, c2, value, font=None, fl=None, align=None, border=False, r2=None):
        r2 = r2 or r
        if c2 > c1 or r2 > r:
            ws.merge_cells(start_row=r, start_column=c1, end_row=r2, end_column=c2)
        cell = ws.cell(r, c1, value)
        if font: cell.font = font
        if fl: cell.fill = fl
        cell.alignment = align or center
        if border:
            for rr in range(r, r2 + 1):
                for c in range(c1, c2 + 1):
                    ws.cell(rr, c).border = brd
                    if fl: ws.cell(rr, c).fill = fl
        return cell


    rn = 1
    try:
        from openpyxl.drawing.image import Image as _XLImage
        logo = _XLImage(logo_path); logo.height = 46; logo.width = 46
        ws.row_dimensions[rn].height = 36
        ws.add_image(logo, f'A{rn}')
    except Exception:
        pass
    rn += 1
    for i, text in enumerate(_title_lines(d)):
        merged(rn, 1, NC, text, Font(bold=True, size=12 if i == 0 else 10)); rn += 1
    merged(rn, 1, NC, ta.CAMPUS, Font(bold=True, size=12, color='FFFFFF'), MAROON)
    ws.row_dimensions[rn].height = 18; rn += 2
    for a, b, c, e in ta._info_rows(d):
        merged(rn, 1, 1, a, Font(bold=True, size=9), SHADE, left, True)
        merged(rn, 2, 4, b, Font(size=9), None, left, True)
        merged(rn, 5, 6, c, Font(bold=True, size=9), SHADE, left, True)
        merged(rn, 7, 9, e, Font(size=9), None, left, True)
        rn += 1
    rn += 1


    # 1. Timetable: A = time, B..H = MON..SUN
    ws.cell(rn, 1, 'WEEKLY TIMETABLE').font = Font(bold=True, size=10); rn += 1
    c = ws.cell(rn, 1, 'TIME'); c.font = Font(bold=True, size=8); c.border = brd; c.alignment = center
    for i, h in enumerate(_day_headers(d)):
        c = ws.cell(rn, 2 + i, h); c.font = Font(bold=True, size=8); c.border = brd; c.alignment = center
    ws.row_dimensions[rn].height = 26; rn += 1
    top = rn
    for k, m in enumerate(SLOTS):
        c = ws.cell(top + k, 1, _slot_label(m)); c.font = Font(bold=True, size=7); c.border = brd; c.alignment = center
        for i in range(7):
            ws.cell(top + k, 2 + i).border = brd
        ws.row_dimensions[top + k].height = 15
    ins = d.get('week_in_semester') or [True] * 7
    for i, ok in enumerate(ins):
        if not ok:
            merged(top, 2 + i, 2 + i, 'No classes', Font(italic=True, size=8, color='555555'),
                   fill(OFF_FILL), center, True, r2=top + len(SLOTS) - 1)
    for i, blocks in grid_model(d).items():
        for r0, r1, text, kind in blocks:
            merged(top + r0, 2 + i, 2 + i, text, Font(bold=True, size=7),
                   fill(MAKEUP_FILL if kind == 'makeup' else CLASS_FILL), center, True, r2=top + r1 - 1)
    rn = top + len(SLOTS) + 1


    # 2. Class list
    ws.cell(rn, 1, 'CLASS LIST').font = Font(bold=True, size=10); rn += 1
    for ci, h in enumerate(LIST_HEADERS, 1):
        c = ws.cell(rn, ci, h); c.font = Font(bold=True, size=8); c.border = brd; c.alignment = center
    rn += 1
    rows = list_rows(d) or [['No classes this week'] + [''] * 8]
    for r in rows:
        for ci, v in enumerate(r):
            c = ws.cell(rn, ci + 1, ta._xl_num(v) if ci == 7 and v else v)
            c.font = Font(size=8, bold=(ci == 8 and v != '')); c.border = brd
            c.alignment = center if ci in LIST_CENTER else left
        rn += 1
    rn += 1


    # 3. Hours per day
    ws.cell(rn, 1, 'TEACHING HOURS PER DAY').font = Font(bold=True, size=10); rn += 1
    heads = [''] + _day_headers(d) + ['TOTAL']
    for ci, h in enumerate(heads, 1):
        c = ws.cell(rn, ci, h); c.font = Font(bold=True, size=8); c.border = brd; c.alignment = center
    ws.row_dimensions[rn].height = 26; rn += 1
    vals, total = hours_row(d)
    for ci, v in enumerate(['HOURS'] + vals + [total], 1):
        c = ws.cell(rn, ci, ta._xl_num(v) if v not in ('', '—', 'HOURS') else v)
        c.font = Font(size=8, bold=(ci in (1, 9))); c.border = brd; c.alignment = center
    rn += 2
    for n in notes(d):
        merged(rn, 1, NC, n, Font(italic=True, size=8), None, left); rn += 1


    buf = io.BytesIO()
    wb.save(buf); buf.seek(0)
    return buf.read()




# ─────────────────────────────────────────────────────────────
#  DOCX
# ─────────────────────────────────────────────────────────────
def gen_docx(d, logo_path=None):
    from docx import Document
    from docx.shared import Pt, Cm, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
    from docx.enum.section import WD_ORIENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn


    BLACK, WHITE = RGBColor(0, 0, 0), RGBColor(0xFF, 0xFF, 0xFF)
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = sec.page_height, sec.page_width
    sec.left_margin = sec.right_margin = Cm(1.2)
    sec.top_margin = sec.bottom_margin = Cm(1.0)
    doc.styles['Normal'].paragraph_format.space_after = Pt(0)
    FULL = 27.3


    def _bg(cell, hex6):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), hex6)
        tcPr.append(shd)


    def _cell(cell, text, bold=False, sz=7, fl=None, align=WD_ALIGN_PARAGRAPH.CENTER, color=None, italic=False):
        cell.text = ''
        p = cell.paragraphs[0]
        run = p.add_run(str(text) if text is not None else '')
        run.font.bold = bold; run.font.italic = italic; run.font.size = Pt(sz)
        run.font.color.rgb = color or BLACK
        p.alignment = align; p.paragraph_format.space_after = Pt(0)
        if fl:
            _bg(cell, fl)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


    def _widths(tbl, cms):
        tbl.autofit = False
        for ci, w in enumerate(cms):
            tbl.columns[ci].width = Cm(w)
            for row in tbl.rows:
                row.cells[ci].width = Cm(w)


    def _para(text, bold=True, sz=9, align=WD_ALIGN_PARAGRAPH.LEFT, before=6, italic=False):
        p = doc.add_paragraph(); p.alignment = align
        p.paragraph_format.space_before = Pt(before); p.paragraph_format.space_after = Pt(2)
        r = p.add_run(text); r.font.bold = bold; r.font.italic = italic; r.font.size = Pt(sz)
        r.font.color.rgb = BLACK


    try:
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(logo_path, width=Cm(1.6))
    except Exception:
        pass
    for i, text in enumerate(_title_lines(d)):
        _para(text, sz=14 if i == 0 else 10, align=WD_ALIGN_PARAGRAPH.CENTER, before=0)
    camp = doc.add_table(rows=1, cols=1); camp.alignment = WD_TABLE_ALIGNMENT.CENTER
    _cell(camp.rows[0].cells[0], ta.CAMPUS, bold=True, sz=11, fl=ta.MAROON_HEX, color=WHITE)
    _widths(camp, [FULL])
    _para('', before=2)
    info = doc.add_table(rows=3, cols=4); info.style = 'Table Grid'; info.alignment = WD_TABLE_ALIGNMENT.CENTER
    for ri, vals in enumerate(ta._info_rows(d)):
        for ci, v in enumerate(vals):
            _cell(info.rows[ri].cells[ci], v, bold=(ci % 2 == 0), sz=8,
                  fl='F2F2F2' if ci % 2 == 0 else None, align=WD_ALIGN_PARAGRAPH.LEFT)
    _widths(info, [2.6, FULL / 2 - 2.6, 2.6, FULL / 2 - 2.6])


    # 1. Timetable
    _para('WEEKLY TIMETABLE', sz=9, before=8)
    grid = doc.add_table(rows=1 + len(SLOTS), cols=8); grid.style = 'Table Grid'
    grid.alignment = WD_TABLE_ALIGNMENT.CENTER
    _widths(grid, [2.1] + [(FULL - 2.1) / 7] * 7)
    _cell(grid.rows[0].cells[0], 'TIME', bold=True, sz=7)
    for i, h in enumerate(_day_headers(d)):
        _cell(grid.rows[0].cells[i + 1], h, bold=True, sz=7)
    for k, m in enumerate(SLOTS):
        _cell(grid.rows[k + 1].cells[0], _slot_label(m), bold=True, sz=6)
        grid.rows[k + 1].height = Cm(0.4)
    ins = d.get('week_in_semester') or [True] * 7
    for i, ok in enumerate(ins):
        if not ok:
            col = grid.rows[1].cells[i + 1].merge(grid.rows[len(SLOTS)].cells[i + 1])
            _cell(col, 'No classes', sz=7, fl=OFF_FILL, italic=True)
    for i, blocks in grid_model(d).items():
        for r0, r1, text, kind in blocks:
            c = grid.rows[r0 + 1].cells[i + 1]
            if r1 - 1 > r0:
                c = c.merge(grid.rows[r1].cells[i + 1])
            _cell(c, text, bold=True, sz=6.5, fl=MAKEUP_FILL if kind == 'makeup' else CLASS_FILL)


    # 2. Class list
    _para('CLASS LIST', sz=9, before=10)
    rows = list_rows(d) or [['No classes this week'] + [''] * 8]
    lt = doc.add_table(rows=1 + len(rows), cols=len(LIST_HEADERS)); lt.style = 'Table Grid'
    lt.alignment = WD_TABLE_ALIGNMENT.CENTER
    _widths(lt, [2.4, 2.2, 4.0, 2.4, 6.3, 2.8, 2.2, 1.6, 3.4])
    for ci, h in enumerate(LIST_HEADERS):
        _cell(lt.rows[0].cells[ci], h, bold=True, sz=7)
    for ri, r in enumerate(rows):
        for ci, v in enumerate(r):
            _cell(lt.rows[ri + 1].cells[ci], v, sz=7, bold=(ci == 8 and v != ''),
                  align=WD_ALIGN_PARAGRAPH.CENTER if ci in LIST_CENTER else WD_ALIGN_PARAGRAPH.LEFT)


    # 3. Hours per day
    _para('TEACHING HOURS PER DAY', sz=9, before=10)
    ht = doc.add_table(rows=2, cols=9); ht.style = 'Table Grid'; ht.alignment = WD_TABLE_ALIGNMENT.CENTER
    _widths(ht, [2.6] + [2.9] * 7 + [4.4])
    for ci, h in enumerate([''] + _day_headers(d) + ['TOTAL']):
        _cell(ht.rows[0].cells[ci], h, bold=True, sz=7)
    vals, total = hours_row(d)
    for ci, v in enumerate(['HOURS'] + vals + [total]):
        _cell(ht.rows[1].cells[ci], v, bold=(ci in (0, 8)), sz=7)
    for n in notes(d):
        _para(n, bold=False, italic=True, sz=7.5, before=4)


    buf = io.BytesIO()
    doc.save(buf); buf.seek(0)
    return buf.read()




# ─────────────────────────────────────────────────────────────
#  PDF
# ─────────────────────────────────────────────────────────────
def gen_pdf(d, logo_path=None):
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib import colors
    from reportlab.lib.units import cm
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import (SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer,
                                    KeepTogether, PageBreak, Image as _RLImage)
    from reportlab.lib.enums import TA_CENTER
    from xml.sax.saxutils import escape as _esc


    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=1.2*cm, rightMargin=1.2*cm,
                            topMargin=0.8*cm, bottomMargin=0.8*cm,
                            title='Weekly Schedule', author='PUP Lopez Scheduling System')
    BLACK, WHITE = colors.black, colors.white
    hx = lambda h: colors.HexColor('#' + h)
    ss = getSampleStyleSheet()
    h1 = ParagraphStyle('H1', parent=ss['Heading1'], textColor=BLACK, fontSize=14, spaceAfter=0,
                        spaceBefore=0, alignment=TA_CENTER, leading=17)
    hsub = ParagraphStyle('HS', parent=ss['Normal'], textColor=BLACK, fontSize=9.5, leading=11.5,
                          alignment=TA_CENTER, fontName='Helvetica-Bold')
    campus = ParagraphStyle('C', parent=ss['Normal'], textColor=WHITE, fontSize=10.5,
                            alignment=TA_CENTER, fontName='Helvetica-Bold')
    sec = ParagraphStyle('S', parent=ss['Normal'], fontName='Helvetica-Bold', fontSize=9,
                         spaceBefore=5, spaceAfter=2, textColor=BLACK)
    cell_c = ParagraphStyle('CC', parent=ss['Normal'], fontName='Helvetica', fontSize=6.8, leading=7.8,
                            alignment=TA_CENTER, textColor=BLACK)
    cell_l = ParagraphStyle('CL', parent=cell_c, alignment=0)
    hdr = ParagraphStyle('HD', parent=cell_c, fontName='Helvetica-Bold', fontSize=6.8)
    blk = ParagraphStyle('BK', parent=cell_c, fontName='Helvetica-Bold', fontSize=6.3, leading=7.2)
    bold_l = ParagraphStyle('BL', parent=cell_l, fontName='Helvetica-Bold', fontSize=7.2)
    note_s = ParagraphStyle('N', parent=cell_l, fontName='Helvetica-Oblique', fontSize=7.5, spaceBefore=3)


    def P(v, st):
        txt = _esc(str(v)).replace('\n', '<br/>') if v not in (None, '') else '&nbsp;'
        return Paragraph(txt, st)


    FULL = landscape(A4)[0] - 2.4*cm
    base = [('GRID', (0, 0), (-1, -1), 0.5, BLACK), ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('LEFTPADDING', (0, 0), (-1, -1), 2), ('RIGHTPADDING', (0, 0), (-1, -1), 2),
            ('TOPPADDING', (0, 0), (-1, -1), 1.5), ('BOTTOMPADDING', (0, 0), (-1, -1), 1.5)]
    story = []


    head = []
    try:
        logo = _RLImage(logo_path, width=1.15*cm, height=1.15*cm); logo.hAlign = 'CENTER'
        head.append(logo)
    except Exception:
        pass
    lines = _title_lines(d)
    head.append(Paragraph(_esc(lines[0]), h1))
    head += [Paragraph(_esc(x), hsub) for x in lines[1:]]
    camp = Table([[Paragraph(ta.CAMPUS, campus)]], colWidths=[FULL])
    camp.setStyle(TableStyle([('BACKGROUND', (0, 0), (-1, -1), hx(ta.MAROON_HEX)),
                              ('BOX', (0, 0), (-1, -1), 0.5, BLACK)]))
    head += [Spacer(1, 0.1*cm), camp, Spacer(1, 0.18*cm)]
    info = Table([[P(a, bold_l), P(b, cell_l), P(c, bold_l), P(e, cell_l)] for a, b, c, e in ta._info_rows(d)],
                 colWidths=[2.6*cm, FULL/2 - 2.6*cm, 2.6*cm, FULL/2 - 2.6*cm])
    info.setStyle(TableStyle(base + [('BACKGROUND', (0, 0), (0, -1), hx('F2F2F2')),
                                     ('BACKGROUND', (2, 0), (2, -1), hx('F2F2F2'))]))
    head.append(info)
    story += head


    # 1. Timetable (fits the rest of page 1)
    data = [[P('TIME', hdr)] + [P(h, hdr) for h in _day_headers(d)]]
    for m in SLOTS:
        data.append([P(_slot_label(m), hdr)] + [''] * 7)
    style = list(base)
    ins = d.get('week_in_semester') or [True] * 7
    for i, ok in enumerate(ins):
        if not ok:
            data[1][i + 1] = P('No classes', ParagraphStyle('OFF', parent=cell_c, fontName='Helvetica-Oblique'))
            style += [('SPAN', (i + 1, 1), (i + 1, len(SLOTS))), ('BACKGROUND', (i + 1, 1), (i + 1, len(SLOTS)), hx(OFF_FILL))]
    for i, blocks in grid_model(d).items():
        for r0, r1, text, kind in blocks:
            data[r0 + 1][i + 1] = P(text, blk)
            style += [('SPAN', (i + 1, r0 + 1), (i + 1, r1)),
                      ('BACKGROUND', (i + 1, r0 + 1), (i + 1, r1), hx(MAKEUP_FILL if kind == 'makeup' else CLASS_FILL))]
    day_w = (FULL - 2.0*cm) / 7
    tt = Table(data, colWidths=[2.0*cm] + [day_w] * 7, rowHeights=[0.62*cm] + [0.36*cm] * len(SLOTS))
    tt.setStyle(TableStyle(style))
    story += [Paragraph('WEEKLY TIMETABLE', sec), tt]


    # 2 + 3 on the next page: class list, hours per day, notes
    story.append(PageBreak())
    rows = list_rows(d) or [['No classes this week'] + [''] * 8]
    lt_data = [[P(h, hdr) for h in LIST_HEADERS]]
    for r in rows:
        lt_data.append([P(v, (bold_l if ci == 8 and v else (cell_c if ci in LIST_CENTER else cell_l)))
                        for ci, v in enumerate(r)])
    lt = Table(lt_data, colWidths=[w*cm for w in (2.4, 2.2, 4.0, 2.4, 6.4, 2.8, 2.2, 1.5, 3.4)], repeatRows=1)
    lt.setStyle(TableStyle(base))
    story += [Paragraph('CLASS LIST', sec), lt]
    vals, total = hours_row(d)
    hd = [[P('', hdr)] + [P(h, hdr) for h in _day_headers(d)] + [P('TOTAL', hdr)],
          [P('HOURS', bold_l)] + [P(v, cell_c) for v in vals] + [P(total, hdr)]]
    ht = Table(hd, colWidths=[2.6*cm] + [2.9*cm] * 7 + [FULL - 2.6*cm - 7 * 2.9*cm])
    ht.setStyle(TableStyle(base))
    story.append(KeepTogether([Paragraph('TEACHING HOURS PER DAY', sec), ht]))
    for n in notes(d):
        story.append(Paragraph(_esc(n), note_s))


    doc.build(story)
    buf.seek(0)
    return buf.read()

GENERATORS = {'xlsx': gen_xlsx, 'csv': gen_csv, 'pdf': gen_pdf, 'docx': gen_docx}