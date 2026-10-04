"""
Faculty side → Teaching Assignment export (CSV / XLSX / PDF / DOCX).

LAYOUT follows the official PUP "Faculty Assignment" form (reference supplied
by the office): title + semester line, an EMP NO / EMP NAME / EMP STATUS ×
COLLEGE / DEPT CODE / DEPARTMENT block, REGULAR LOAD and PART-TIME tables
(Subject Code, Subject Description, Units, Year & Section, Subj. Ref., Time,
Time Code, Day/s, Room, Effectivity) with blank ruled rows, each followed by
its total, then TEACHING LOAD PER DAY (HOURS), OFFICIAL TIME / ADVISING TIME,
the subject-reference legend and a signature line.

DESIGN stays the system's own, identical to its other exports (Reports >
Class Schedule, Faculty List…): centred PUP logo, bold black title, the maroon
"LOPEZ, QUEZON CAMPUS" banner, black-grid tables with bold headers. The logo is
the shared file (app._SCH_LOGO_PATH, passed in by the caller).

DATA is exactly the Teaching Assignment tab's payload
(app._my_teaching_load_payload). Fields the system does not record are left
blank for hand-filling, as on the paper form: TIME CODE, OFFICIAL TIME and
ADVISING TIME. SUBJ. REF. is derived from the program: diploma programs
(codes starting with "D", e.g. DIT) → (T) iTech, all others → (C) College.
"""
import csv
import io
import re

SEM_LABELS = {'A': '1st Semester', 'B': '2nd Semester', 'C': 'Summer'}
_SEM_WORDS = {'1st Semester': 'FIRST SEMESTER', '2nd Semester': 'SECOND SEMESTER', 'Summer': 'SUMMER'}

LOAD_HEADERS = ['SUBJECT CODE', 'SUBJECT DESCRIPTION', 'UNITS', 'YEAR & SECTION', 'SUBJ. REF.',
                'TIME', 'TIME CODE', 'DAY/S', 'ROOM', 'EFFTVTY.']
LOAD_CENTER = {0, 2, 3, 4, 5, 6, 7, 8, 9}          # everything but the description
LOAD_COL_CM = [1.75, 3.95, 0.95, 2.05, 0.95, 3.45, 1.15, 1.2, 1.4, 1.75]   # = 18.6 cm (A4 portrait, 1.2 cm margins)
LOAD_XLSX_W = [12, 28, 7, 14, 7, 26, 8, 8, 10, 12]
MIN_LOAD_ROWS = 10                                   # ruled rows per table, like the paper form

DAYS = ['Monday', 'Tuesday', 'Wednesday', 'Thursday', 'Friday', 'Saturday', 'Sunday']
DAY_HEADERS = ['MON', 'TUE', 'WED', 'THUR', 'FRI', 'SAT', 'SUN']
LEGEND = ('SUBJECT REFERENCE LEGEND: (C) - College, (OU) - Open University, (GS) - Graduate School, '
          '(PB) - Post Bac, (L) - Law, (T) - iTech')

COLLEGE, DEPARTMENT = 'LOPEZ, QUEZON', 'LOPEZ BRANCH'
MAROON_HEX = '7A0100'
CAMPUS = 'LOPEZ, QUEZON CAMPUS'

MIME = {
    'xlsx': ('application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', '.xlsx'),
    'csv':  ('text/csv', '.csv'),
    'pdf':  ('application/pdf', '.pdf'),
    'docx': ('application/vnd.openxmlformats-officedocument.wordprocessingml.document', '.docx'),
}


# ─────────────────────────────────────────────────────────────
#  Shared content (one model, rendered by every format)
# ─────────────────────────────────────────────────────────────
def clean_filename(name):
    """Safe download name without extension (keeps letters, digits, - _ .)."""
    name = (name or '').strip()
    for ext in ('.xlsx', '.csv', '.pdf', '.docx', '.zip'):
        if name.lower().endswith(ext):
            name = name[:-len(ext)]
    name = re.sub(r'\s+', '_', name)
    name = re.sub(r'[^A-Za-z0-9._-]', '', name)
    return name.strip('._')[:150]


def export_filename(fullname, sem_label, ay_label):
    """Teaching_Assignment_[FacultyName]_[Semester]_[AcademicYear]"""
    return clean_filename(f"Teaching_Assignment_{fullname}_{sem_label}_{ay_label}") or 'Teaching_Assignment'


def _num(v):
    if v is None or v == '':
        return ''
    try:
        f = float(v)
    except (TypeError, ValueError):
        return str(v)
    return str(int(f)) if f == int(f) else f'{f:g}'


def _xl_num(v):
    """Numeric cell value for Excel (so units can be summed); text fallback."""
    try:
        f = float(v)
    except (TypeError, ValueError):
        return v if v not in (None, '') else ''
    return int(f) if f == int(f) else f


def _short_time(t):
    """'07:30 AM' → '7:30AM', '12:00 PM' → '12PM' (the form's compact style)."""
    m = re.match(r'\s*0?(\d{1,2}):(\d{2})\s*([AP]M)', t or '', re.I)
    if not m:
        return (t or '').strip()
    h, mm, ap = m.groups()
    return f"{int(h)}{'' if mm == '00' else ':' + mm}{ap.upper()}"


def compact_time_range(time_range):
    """'07:30 AM - 12:00 PM, 01:00 PM - 02:30 PM' → '7:30AM-12PM/1PM-2:30PM'."""
    parts = []
    for rng in str(time_range or '').split(','):
        ends = [e for e in rng.split(' - ')]
        if len(ends) == 2:
            parts.append(f"{_short_time(ends[0])}-{_short_time(ends[1])}")
        elif rng.strip() and rng.strip() != '—':
            parts.append(rng.strip())
    return '/'.join(parts)


def subject_reference(year_section):
    """(T) iTech for diploma programs (codes starting with D, e.g. DIT); (C) College otherwise."""
    m = re.match(r'\s*([A-Za-z]+)', year_section or '')
    if not m:
        return ''
    return 'T' if m.group(1).upper().startswith('D') else 'C'


def _title_lines(d):
    sem = _SEM_WORDS.get(d.get('sem_label') or '', (d.get('sem_label') or '').upper())
    line2 = f"{sem}, SY {d['ay_label']}" if d.get('ay_label') else sem
    lines = ['FACULTY ASSIGNMENT', line2]
    if d.get('source') == 'local':
        lines.append('LOCAL SCHEDULE')
    return lines


def _info_rows(d):
    """3 rows × (label, value, label, value) — the form's header block."""
    return [
        ('EMP NO', d.get('employeenumber') or '', 'COLLEGE', COLLEGE),
        ('EMP NAME', (d.get('fullname') or '').upper(), 'DEPT CODE', d.get('dept_code') or ''),
        ('EMP STATUS', d.get('employee_status') or d.get('employee_type') or '', 'DEPARTMENT', DEPARTMENT),
    ]


def _load_row(r):
    return [r.get('subjectcode') or '', r.get('subjectname') or '', _num(r.get('units')),
            r.get('year_section') or '', subject_reference(r.get('year_section')),
            compact_time_range(r.get('time_range')), '',             # TIME CODE: not recorded
            r.get('days') or '', r.get('room') or '', r.get('effectivity') or '']


def _sections(d):
    t = d.get('totals') or {}
    out = [('REGULAR LOAD', d.get('regular') or [], 'Total REGULAR LOAD', t.get('regularUnits')),
           ('PART-TIME', d.get('partTime') or [], 'Total PART-TIME', t.get('partTimeUnits'))]
    return out


def _ts_line(d):
    ts = (d.get('totals') or {}).get('tsUnits')
    return ('Total TEACHING SUBSTITUTION (TS)', ts) if ts and float(ts) > 0 else None


def _per_day_rows(d):
    pd = d.get('per_day_hours') or {}
    reg, pt = pd.get('regular') or {}, pd.get('part_time') or {}
    fmt = lambda v: _num(v) if v else ''
    return [
        ['REGULAR']   + [fmt(reg.get(day)) for day in DAYS],
        ['PART-TIME'] + [fmt(pt.get(day)) for day in DAYS],
        ['TOTAL']     + [fmt((reg.get(day) or 0) + (pt.get(day) or 0)) for day in DAYS],
    ]


def _official_rows():
    return [['OFFICIAL TIME'] + [''] * 7, ['ADVISING TIME'] + [''] * 7]


# ─────────────────────────────────────────────────────────────
#  CSV — same sections, text letterhead (a CSV can't hold the logo)
# ─────────────────────────────────────────────────────────────
def gen_csv(d, logo_path=None):
    out = io.StringIO()
    w = csv.writer(out)
    for line in _title_lines(d):
        w.writerow([line])
    w.writerow([CAMPUS])
    w.writerow([])
    for a, b, c, e in _info_rows(d):
        w.writerow([a, b, c, e])
    for title, rows, total_lbl, total in _sections(d):
        w.writerow([])
        w.writerow([title])
        w.writerow(LOAD_HEADERS)
        for r in rows:
            w.writerow(_load_row(r))
        w.writerow([total_lbl, _num(total)])
    if _ts_line(d):
        w.writerow([_ts_line(d)[0], _num(_ts_line(d)[1])])
    w.writerow([])
    w.writerow(['TEACHING LOAD PER DAY (HOURS)'])
    w.writerow([''] + DAY_HEADERS)
    for row in _per_day_rows(d):
        w.writerow(row)
    w.writerow([])
    w.writerow(['OFFICIAL TIME / ADVISING TIME'])
    w.writerow([''] + DAY_HEADERS)
    for row in _official_rows():
        w.writerow(row)
    w.writerow([])
    w.writerow([LEGEND])
    return out.getvalue().encode('utf-8-sig')


# ─────────────────────────────────────────────────────────────
#  XLSX
# ─────────────────────────────────────────────────────────────
def gen_xlsx(d, logo_path=None):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = 'Faculty Assignment'
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = 'portrait'
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0

    MAROON = PatternFill('solid', fgColor=MAROON_HEX)
    SHADE = PatternFill('solid', fgColor='F2F2F2')
    thin = Side(style='thin', color='000000')
    brd = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    left = Alignment(horizontal='left', vertical='center', wrap_text=True)
    N = len(LOAD_HEADERS)
    for ci, wdt in enumerate(LOAD_XLSX_W, 1):
        ws.column_dimensions[get_column_letter(ci)].width = wdt

    def merged(r, c1, c2, value, font=None, fill=None, align=None, border=False):
        if c2 > c1:
            ws.merge_cells(start_row=r, start_column=c1, end_row=r, end_column=c2)
        cell = ws.cell(r, c1, value)
        if font: cell.font = font
        if fill: cell.fill = fill
        cell.alignment = align or center
        if border:
            for c in range(c1, c2 + 1):
                ws.cell(r, c).border = brd
        return cell

    rn = 1
    try:
        from openpyxl.drawing.image import Image as _XLImage
        logo = _XLImage(logo_path)
        logo.height = 46; logo.width = 46
        ws.row_dimensions[rn].height = 36
        ws.add_image(logo, f'A{rn}')
    except Exception:
        pass
    rn += 1
    for i, text in enumerate(_title_lines(d)):
        merged(rn, 1, N, text, Font(bold=True, size=12 if i == 0 else 10))
        ws.row_dimensions[rn].height = 18; rn += 1
    merged(rn, 1, N, CAMPUS, Font(bold=True, size=12, color='FFFFFF'), MAROON)
    ws.row_dimensions[rn].height = 18; rn += 2

    # Header block: label A-B | value C-E | label F-G | value H-J
    for a, b, c, e in _info_rows(d):
        merged(rn, 1, 2, a, Font(bold=True, size=9), SHADE, left, True)
        merged(rn, 3, 5, b, Font(size=9), None, left, True)
        merged(rn, 6, 7, c, Font(bold=True, size=9), SHADE, left, True)
        merged(rn, 8, 10, e, Font(size=9), None, left, True)
        rn += 1
    rn += 1

    for title, rows, total_lbl, total in _sections(d):
        ws.cell(rn, 1, title).font = Font(bold=True, size=10)
        rn += 1
        for ci, h in enumerate(LOAD_HEADERS, 1):
            c = ws.cell(rn, ci, h); c.font = Font(bold=True, size=8); c.border = brd; c.alignment = center
        ws.row_dimensions[rn].height = 24; rn += 1
        for i in range(max(len(rows), MIN_LOAD_ROWS)):
            vals = _load_row(rows[i]) if i < len(rows) else [''] * N
            for ci, v in enumerate(vals):
                c = ws.cell(rn, ci + 1, _xl_num(rows[i].get('units')) if (ci == 2 and i < len(rows)) else v)
                c.font = Font(size=8); c.border = brd
                c.alignment = center if ci in LOAD_CENTER else left
            rn += 1
        ws.cell(rn, 1, total_lbl).font = Font(bold=True, size=9)
        c = ws.cell(rn, 3, _xl_num(total)); c.font = Font(bold=True, size=9); c.alignment = center
        rn += 2
    ts = _ts_line(d)
    if ts:
        ws.cell(rn, 1, ts[0]).font = Font(bold=True, size=9)
        ws.cell(rn, 3, _xl_num(ts[1])).font = Font(bold=True, size=9)
        rn += 2

    # Day grids: label A-B | MON..SUN in C..I
    for grid_title, grid_rows, height in (('TEACHING LOAD PER DAY (HOURS)', _per_day_rows(d), 16),
                                          ('OFFICIAL TIME / ADVISING TIME', _official_rows(), 30)):
        merged(rn, 1, 9, grid_title, Font(bold=True, size=10))
        rn += 1
        merged(rn, 1, 2, '', border=True)
        for i, h in enumerate(DAY_HEADERS):
            c = ws.cell(rn, 3 + i, h); c.font = Font(bold=True, size=8); c.border = brd; c.alignment = center
        rn += 1
        for row in grid_rows:
            merged(rn, 1, 2, row[0], Font(bold=True, size=8), None, left, True)
            for i, v in enumerate(row[1:]):
                c = ws.cell(rn, 3 + i, _xl_num(v) if v != '' else '')
                c.font = Font(bold=(row[0] == 'TOTAL'), size=8); c.border = brd; c.alignment = center
            ws.row_dimensions[rn].height = height
            rn += 1
        rn += 1

    merged(rn, 1, N, LEGEND, Font(bold=True, size=8), None, left)
    rn += 3
    merged(rn, 7, 10, '______________________________', Font(size=9))
    merged(rn + 1, 7, 10, 'Signature over Printed Name', Font(size=8, italic=True))

    buf = io.BytesIO()
    wb.save(buf); buf.seek(0)
    return buf.read()


# ─────────────────────────────────────────────────────────────
#  DOCX
# ─────────────────────────────────────────────────────────────
def gen_docx(d, logo_path=None):
    from docx import Document
    from docx.shared import Pt, Cm, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    BLACK, WHITE = RGBColor(0, 0, 0), RGBColor(0xFF, 0xFF, 0xFF)
    doc = Document()
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(1.2)
        s.top_margin = s.bottom_margin = Cm(1.0)
    doc.styles['Normal'].paragraph_format.space_after = Pt(0)

    def _bg(cell, hex6):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), hex6)
        tcPr.append(shd)

    def _cell(cell, text, bold=False, sz=7, fill=None, align=WD_ALIGN_PARAGRAPH.CENTER, color=None):
        cell.text = ''
        p = cell.paragraphs[0]
        run = p.add_run(str(text) if text is not None else '')
        run.font.bold = bold; run.font.size = Pt(sz); run.font.color.rgb = color or BLACK
        p.alignment = align
        p.paragraph_format.space_after = Pt(0)
        if fill:
            _bg(cell, fill)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    def _widths(tbl, cms):
        tbl.autofit = False
        for ci, w in enumerate(cms):
            tbl.columns[ci].width = Cm(w)
            for row in tbl.rows:
                row.cells[ci].width = Cm(w)

    def _para(text, bold=True, sz=9, align=WD_ALIGN_PARAGRAPH.LEFT, before=6, italic=False):
        p = doc.add_paragraph(); p.alignment = align
        p.paragraph_format.space_before = Pt(before); p.paragraph_format.space_after = Pt(2)
        r = p.add_run(text); r.font.bold = bold; r.font.italic = italic; r.font.size = Pt(sz)
        r.font.color.rgb = BLACK
        return p

    try:
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(logo_path, width=Cm(1.8))
    except Exception:
        pass
    for i, text in enumerate(_title_lines(d)):
        _para(text, sz=14 if i == 0 else 10, align=WD_ALIGN_PARAGRAPH.CENTER, before=0)
    camp = doc.add_table(rows=1, cols=1)
    camp.alignment = WD_TABLE_ALIGNMENT.CENTER
    _cell(camp.rows[0].cells[0], CAMPUS, bold=True, sz=11, fill=MAROON_HEX, color=WHITE)
    _widths(camp, [sum(LOAD_COL_CM)])
    _para('', before=2)

    info = doc.add_table(rows=3, cols=4)
    info.style = 'Table Grid'; info.alignment = WD_TABLE_ALIGNMENT.CENTER
    for ri, vals in enumerate(_info_rows(d)):
        cells = info.rows[ri].cells
        for ci, v in enumerate(vals):
            _cell(cells[ci], v, bold=(ci % 2 == 0), sz=8, fill='F2F2F2' if ci % 2 == 0 else None,
                  align=WD_ALIGN_PARAGRAPH.LEFT)
    _widths(info, [2.4, 6.9, 2.4, 6.9])

    for title, rows, total_lbl, total in _sections(d):
        _para(title, sz=9, before=8)
        n_rows = max(len(rows), MIN_LOAD_ROWS)
        tbl = doc.add_table(rows=1 + n_rows, cols=len(LOAD_HEADERS))
        tbl.style = 'Table Grid'; tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        _widths(tbl, LOAD_COL_CM)
        for ci, h in enumerate(LOAD_HEADERS):
            _cell(tbl.rows[0].cells[ci], h, bold=True, sz=6.5)
        for i in range(n_rows):
            vals = _load_row(rows[i]) if i < len(rows) else [''] * len(LOAD_HEADERS)
            cells = tbl.rows[i + 1].cells
            for ci, v in enumerate(vals):
                _cell(cells[ci], v, sz=7,
                      align=WD_ALIGN_PARAGRAPH.CENTER if ci in LOAD_CENTER else WD_ALIGN_PARAGRAPH.LEFT)
        _para(f"{total_lbl}   {_num(total)}", sz=8, before=2)
    ts = _ts_line(d)
    if ts:
        _para(f"{ts[0]}   {_num(ts[1])}", sz=8, before=2)

    grid_w = [3.0] + [2.2] * 7
    for grid_title, grid_rows in (('TEACHING LOAD PER DAY (HOURS)', _per_day_rows(d)),
                                  ('OFFICIAL TIME / ADVISING TIME', _official_rows())):
        _para(grid_title, sz=9, align=WD_ALIGN_PARAGRAPH.CENTER, before=8)
        g = doc.add_table(rows=1 + len(grid_rows), cols=8)
        g.style = 'Table Grid'; g.alignment = WD_TABLE_ALIGNMENT.CENTER
        _widths(g, grid_w)
        _cell(g.rows[0].cells[0], '', sz=7)
        for i, h in enumerate(DAY_HEADERS):
            _cell(g.rows[0].cells[i + 1], h, bold=True, sz=7)
        for ri, row in enumerate(grid_rows):
            cells = g.rows[ri + 1].cells
            _cell(cells[0], row[0], bold=True, sz=7, align=WD_ALIGN_PARAGRAPH.LEFT)
            for i, v in enumerate(row[1:]):
                _cell(cells[i + 1], v, bold=(row[0] == 'TOTAL'), sz=7)
            if grid_title.startswith('OFFICIAL'):
                g.rows[ri + 1].height = Cm(0.9)

    _para(LEGEND, sz=7, before=8)
    _para('______________________________', bold=False, sz=9, align=WD_ALIGN_PARAGRAPH.RIGHT, before=22)
    _para('Signature over Printed Name', bold=False, italic=True, sz=8, align=WD_ALIGN_PARAGRAPH.RIGHT, before=0)

    buf = io.BytesIO()
    doc.save(buf); buf.seek(0)
    return buf.read()


# ─────────────────────────────────────────────────────────────
#  PDF
# ─────────────────────────────────────────────────────────────
def gen_pdf(d, logo_path=None):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib import colors
    from reportlab.lib.units import cm
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import (SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer,
                                    KeepTogether, Image as _RLImage)
    from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT
    from xml.sax.saxutils import escape as _esc

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=1.2*cm, rightMargin=1.2*cm,
                            topMargin=0.9*cm, bottomMargin=0.9*cm,
                            title='Faculty Assignment', author='PUP Lopez Scheduling System')
    BLACK, WHITE = colors.black, colors.white
    MAROON = colors.HexColor('#' + MAROON_HEX)
    SHADE = colors.HexColor('#F2F2F2')
    ss = getSampleStyleSheet()
    h1 = ParagraphStyle('H1', parent=ss['Heading1'], textColor=BLACK, fontSize=14, spaceAfter=1, alignment=TA_CENTER)
    hsub = ParagraphStyle('HSub', parent=ss['Normal'], textColor=BLACK, fontSize=10, spaceAfter=4,
                          alignment=TA_CENTER, fontName='Helvetica-Bold')
    campus = ParagraphStyle('Campus', parent=ss['Normal'], textColor=WHITE, fontSize=11,
                            alignment=TA_CENTER, fontName='Helvetica-Bold')
    sec = ParagraphStyle('Sec', parent=ss['Normal'], fontName='Helvetica-Bold', fontSize=9,
                         spaceBefore=6, spaceAfter=2, textColor=BLACK)
    sec_c = ParagraphStyle('SecC', parent=sec, alignment=TA_CENTER)
    cell_l = ParagraphStyle('CL', parent=ss['Normal'], fontName='Helvetica', fontSize=7, leading=8.4, textColor=BLACK)
    cell_c = ParagraphStyle('CC', parent=cell_l, alignment=TA_CENTER)
    hdr = ParagraphStyle('HD', parent=cell_c, fontName='Helvetica-Bold', fontSize=6.5, leading=7.6)
    bold_l = ParagraphStyle('BL', parent=cell_l, fontName='Helvetica-Bold', fontSize=7.5)
    bold_c = ParagraphStyle('BC', parent=cell_c, fontName='Helvetica-Bold')
    total_s = ParagraphStyle('TOT', parent=cell_l, fontName='Helvetica-Bold', fontSize=8, spaceBefore=2)
    legend_s = ParagraphStyle('LEG', parent=cell_l, fontName='Helvetica-Bold', fontSize=7, spaceBefore=8)
    sig_s = ParagraphStyle('SIG', parent=cell_l, fontSize=8, alignment=TA_RIGHT)

    def P(v, st):
        return Paragraph(_esc(str(v)) if v not in (None, '') else '&nbsp;', st)

    FULL = sum(LOAD_COL_CM) * cm
    grid = [('GRID', (0, 0), (-1, -1), 0.5, BLACK), ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('LEFTPADDING', (0, 0), (-1, -1), 2.5), ('RIGHTPADDING', (0, 0), (-1, -1), 2.5),
            ('TOPPADDING', (0, 0), (-1, -1), 2), ('BOTTOMPADDING', (0, 0), (-1, -1), 2)]
    story = []

    try:
        logo = _RLImage(logo_path, width=1.35*cm, height=1.35*cm); logo.hAlign = 'CENTER'
        story.append(logo)
    except Exception:
        pass
    lines = _title_lines(d)
    story.append(Paragraph(_esc(lines[0]), h1))
    for extra in lines[1:]:
        story.append(Paragraph(_esc(extra), hsub))
    camp = Table([[Paragraph(CAMPUS, campus)]], colWidths=[FULL])
    camp.setStyle(TableStyle([('BACKGROUND', (0, 0), (-1, -1), MAROON), ('BOX', (0, 0), (-1, -1), 0.5, BLACK)]))
    story += [camp, Spacer(1, 0.25*cm)]

    info = Table([[P(a, bold_l), P(b, cell_l), P(c, bold_l), P(e, cell_l)] for a, b, c, e in _info_rows(d)],
                 colWidths=[2.4*cm, FULL/2 - 2.4*cm, 2.4*cm, FULL/2 - 2.4*cm])
    info.setStyle(TableStyle(grid + [('BACKGROUND', (0, 0), (0, -1), SHADE), ('BACKGROUND', (2, 0), (2, -1), SHADE)]))
    story.append(info)

    for title, rows, total_lbl, total in _sections(d):
        data = [[P(h, hdr) for h in LOAD_HEADERS]]
        for i in range(max(len(rows), MIN_LOAD_ROWS)):
            vals = _load_row(rows[i]) if i < len(rows) else [''] * len(LOAD_HEADERS)
            data.append([P(v, cell_c if ci in LOAD_CENTER else cell_l) for ci, v in enumerate(vals)])
        tbl = Table(data, colWidths=[w*cm for w in LOAD_COL_CM], repeatRows=1,
                    rowHeights=[None] + [0.5*cm if i >= len(rows) else None for i in range(len(data) - 1)])
        tbl.setStyle(TableStyle(grid))
        story.append(KeepTogether([Paragraph(title, sec), tbl,
                                   Paragraph(f"{_esc(total_lbl)}&nbsp;&nbsp;&nbsp;{_esc(_num(total))}", total_s)]))
    ts = _ts_line(d)
    if ts:
        story.append(Paragraph(f"{_esc(ts[0])}&nbsp;&nbsp;&nbsp;{_esc(_num(ts[1]))}", total_s))

    day_w = [3.0*cm] + [(FULL - 3.0*cm) / 7] * 7
    for grid_title, grid_rows, row_h in (('TEACHING LOAD PER DAY (HOURS)', _per_day_rows(d), None),
                                         ('OFFICIAL TIME / ADVISING TIME', _official_rows(), 0.85*cm)):
        data = [[P('', cell_c)] + [P(h, hdr) for h in DAY_HEADERS]]
        for row in grid_rows:
            st = bold_c if row[0] == 'TOTAL' else cell_c
            data.append([P(row[0], bold_l)] + [P(v, st) for v in row[1:]])
        t = Table(data, colWidths=day_w, rowHeights=[None] + [row_h] * len(grid_rows))
        t.setStyle(TableStyle(grid))
        story.append(KeepTogether([Paragraph(grid_title, sec_c), t]))

    story.append(Paragraph(_esc(LEGEND), legend_s))
    story.append(Spacer(1, 0.9*cm))
    story.append(Paragraph('______________________________', sig_s))
    story.append(Paragraph('<i>Signature over Printed Name</i>', sig_s))

    doc.build(story)
    buf.seek(0)
    return buf.read()


GENERATORS = {'xlsx': gen_xlsx, 'csv': gen_csv, 'pdf': gen_pdf, 'docx': gen_docx}
