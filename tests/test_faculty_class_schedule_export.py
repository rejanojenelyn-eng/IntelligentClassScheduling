"""
Faculty → Class Schedule (SIS) → Export (XLSX / CSV / PDF / DOCX).

The export must be exactly what the Faculty page shows for the same filters
(active term, Published only — never Draft), built with the same official
generators as the Academic Head's Class Schedule export / Reports > Class
Schedule, in one chosen format and layout.
"""
import csv
import io

import pytest

from conftest import requires_db, DB_AVAILABLE

pytestmark = requires_db

if DB_AVAILABLE:
    import app as app_module
    from database import query_db

FAC = '12079'
ROUTE = '/faculty/schedule/export'


def _client(username=FAC, role='Faculty'):
    c = app_module.app.test_client()
    with c.session_transaction() as s:
        s.update(loggedin=True, username=username, role=role,
                 must_change_password=False, account_setup_complete=True)
    return c


@pytest.fixture
def term():
    t = query_db("SELECT academicyearid, semestertype, semesterid FROM semester WHERE isactive LIMIT 1", one=True)
    if not t or not query_db("SELECT 1 FROM accounts WHERE username = %s AND isactive", [FAC]):
        pytest.skip('needs an active term and faculty account 12079')
    return t


def _page_codes(c, term, **filters):
    q = '&'.join(f'{k}={v}' for k, v in filters.items())
    r = c.get(f"/api/faculty/schedule?ay_id={term['academicyearid']}&semester={term['semestertype']}&{q}")
    return sorted({s['subjectcode'] for s in r.get_json()['sessions']})


def _csv_codes(data):
    rows = list(csv.reader(io.StringIO(data.decode('utf-8-sig'))))
    return sorted({r[1] for r in rows[1:] if len(r) > 1})


def _text(fmt, data):
    if fmt == 'xlsx':
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data))
        return '\n'.join(str(c.value) for ws in wb for r in ws.iter_rows() for c in r if c.value is not None)
    if fmt == 'docx':
        from docx import Document
        d = Document(io.BytesIO(data))
        return '\n'.join([p.text for p in d.paragraphs] + [c.text for t in d.tables for r in t.rows for c in r.cells])
    import pdfplumber
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        return '\n'.join(p.extract_text() or '' for p in pdf.pages)


# ── what gets exported ──────────────────────────────────────────────────────

def test_instructor_filter_matches_the_page(term):
    c = _client()
    r = c.post(ROUTE, json={'format': 'csv', 'emp_num': FAC})
    assert r.status_code == 200
    assert _csv_codes(r.data) == _page_codes(c, term, emp_num=FAC)


def test_section_filter_matches_the_page(term):
    c = _client()
    sec = query_db("""SELECT sc.sectionid, pyl.programcode, pyl.yearlevel FROM schedule sc
                      JOIN schedule_version sv ON sv.scheduleid = sc.scheduleid AND sv.status = 'Published'
                      JOIN sections s ON s.sectionid = sc.sectionid
                      JOIN program_yearlevel pyl ON pyl.programyearlevelid = s.programyearlevelid
                      WHERE sc.employeenumber = %s AND sc.semesterid = %s LIMIT 1""",
                   [FAC, term['semesterid']], one=True)
    r = c.post(ROUTE, json={'format': 'csv', 'program': sec['programcode'],
                            'year_level': sec['yearlevel'], 'section_id': sec['sectionid']})
    assert r.status_code == 200
    assert _csv_codes(r.data) == _page_codes(c, term, program=sec['programcode'],
                                             year_level=sec['yearlevel'], section_id=sec['sectionid'])


@pytest.mark.parametrize('fmt', ['xlsx', 'pdf', 'docx'])
@pytest.mark.parametrize('layout', ['table', 'calendar'])
def test_documents_use_the_official_class_schedule_design(term, fmt, layout):
    c = _client()
    r = c.post(ROUTE, json={'format': fmt, 'layout': layout, 'emp_num': FAC, 'filename': 'My Schedule'})
    assert r.status_code == 200, r.get_json()
    assert f'filename="My_Schedule.{fmt}"' in r.headers['Content-Disposition']
    text = _text(fmt, r.data)
    assert 'SUBJECT OFFERINGS FOR' in text and 'LOPEZ, QUEZON CAMPUS' in text
    for code in _page_codes(c, term, emp_num=FAC):
        assert code.replace(' ', '') in text.replace(' ', ''), (fmt, layout, code)


def test_same_bytes_as_the_academic_head_generator_for_the_same_rows(term):
    """The faculty file is produced by the shared _sch_export_bytes pipeline."""
    c = _client()
    ours = c.post(ROUTE, json={'format': 'csv', 'emp_num': FAC}).data
    conn = app_module.get_db_connection()
    try:
        cur = conn.cursor(cursor_factory=app_module.RealDictCursor)
        ctx = app_module._sch_export_context(cur, [term['academicyearid']], [term['semestertype']], [], [],
                                             'table', published_only=True)
        rows, groups, cr, cg = app_module._sch_filter_export(cur, ctx[0], ctx[1], ctx[2], ctx[3], None, FAC)
        theirs = app_module._sch_export_bytes('csv', 'table', rows, groups, cr, cg, ctx[4], ctx[5], ctx[6])
    finally:
        conn.close()
    assert ours == theirs


# ── Draft data never leaks to the Faculty export ────────────────────────────

@pytest.fixture
def draft_offering(term):
    """A Draft-only schedule for faculty 12079 in the active term."""
    sec = query_db("""SELECT sc.sectionid FROM schedule sc WHERE sc.employeenumber = %s AND sc.semesterid = %s
                      LIMIT 1""", [FAC, term['semesterid']], one=True)
    cs = query_db("""SELECT cs.curriculumsubjectid, cs.subjectcode FROM curriculumsubject cs
                     WHERE NOT EXISTS (SELECT 1 FROM schedule s WHERE s.curriculumsubjectid = cs.curriculumsubjectid
                                       AND s.sectionid = %s AND s.semesterid = %s)
                       AND cs.subjectcode NOT IN (SELECT cs2.subjectcode FROM schedule s2
                           JOIN curriculumsubject cs2 ON cs2.curriculumsubjectid = s2.curriculumsubjectid
                           WHERE s2.employeenumber = %s)
                     ORDER BY cs.curriculumsubjectid LIMIT 1""",
                  [sec['sectionid'], term['semesterid'], FAC], one=True)
    sid = query_db("""INSERT INTO schedule (curriculumsubjectid, sectionid, employeenumber, semesterid)
                      VALUES (%s,%s,%s,%s) RETURNING scheduleid""",
                   [cs['curriculumsubjectid'], sec['sectionid'], FAC, term['semesterid']], one=True)['scheduleid']
    vid = query_db("INSERT INTO schedule_version (scheduleid, version_number, status) VALUES (%s,1,'Draft') RETURNING versionid",
                   [sid], one=True)['versionid']
    tid = {r['v']: r['timeid'] for r in query_db("SELECT timeid, TO_CHAR(timevalue,'HH24:MI') v FROM timeslot")}
    query_db("""INSERT INTO schedule_sessions (versionid, daydesc, starttimeid, endtimeid, roomid)
                VALUES (%s,'Sunday',%s,%s,NULL) RETURNING 1""", [vid, tid['08:00'], tid['09:00']])
    try:
        yield cs['subjectcode']
    finally:
        query_db("DELETE FROM schedule WHERE scheduleid = %s RETURNING 1", [sid])


def test_draft_schedules_are_not_exported(term, draft_offering):
    c = _client()
    codes = _csv_codes(c.post(ROUTE, json={'format': 'csv', 'emp_num': FAC}).data)
    assert draft_offering not in codes
    # …while the Academic Head's (Draft-inclusive) fetch still sees it — unchanged behaviour
    conn = app_module.get_db_connection()
    try:
        cur = conn.cursor(cursor_factory=app_module.RealDictCursor)
        all_rows = app_module._sch_exp_fetch(cur, [term['academicyearid']], [term['semestertype']], [], [])
    finally:
        conn.close()
    assert draft_offering in {r['SubjectCode'] for r in all_rows}


# ── validation & access ─────────────────────────────────────────────────────

def test_requires_a_section_or_instructor(term):
    r = _client().post(ROUTE, json={'format': 'pdf', 'program': 'BSIT'})
    assert r.status_code == 400 and 'Section or an Instructor' in r.get_json()['error']


def test_only_one_valid_format(term):
    c = _client()
    assert c.post(ROUTE, json={'format': 'txt', 'emp_num': FAC}).status_code == 400
    assert c.post(ROUTE, json={'format': '', 'emp_num': FAC}).status_code == 400


def test_empty_selection_returns_a_clear_message(term):
    other = query_db("""SELECT f.employeenumber FROM faculty f WHERE NOT EXISTS (
                          SELECT 1 FROM schedule s JOIN schedule_version sv ON sv.scheduleid = s.scheduleid
                          WHERE s.employeenumber = f.employeenumber AND sv.status = 'Published'
                            AND s.semesterid = %s) LIMIT 1""", [term['semesterid']], one=True)
    r = _client().post(ROUTE, json={'format': 'pdf', 'emp_num': other['employeenumber']})
    assert r.status_code == 404 and 'no published schedule' in r.get_json()['error']


def test_faculty_only():
    assert app_module.app.test_client().post(ROUTE, json={'format': 'pdf', 'emp_num': FAC}).status_code == 403
    assert _client('94045', 'Academic Head').post(ROUTE, json={'format': 'pdf', 'emp_num': FAC}).status_code == 403


# ── the shared filter refactor didn't change Reports > Class Schedule ───────

def test_reports_class_schedule_instructor_filter_still_works(term):
    c = _client('94045', 'Academic Head')
    r = c.post('/reports/export/class_schedule', json={
        'formats': ['csv'], 'ay_ids': [term['academicyearid']], 'sem_types': [term['semestertype']],
        'faculty': FAC})
    assert r.status_code == 200
    names = {row[0] for row in csv.reader(io.StringIO(r.data.decode('utf-8-sig')))}
    assert names - {'Instructor'} and all(n.startswith('ZURBANO') for n in names - {'Instructor'})
