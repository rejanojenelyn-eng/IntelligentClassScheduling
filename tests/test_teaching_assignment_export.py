"""
Faculty side → Teaching Assignment → Export (XLSX / CSV / PDF / DOCX).

Checks each format downloads with the Teaching_Assignment_[Faculty]_[Semester]_[AY]
name, carries the official letterhead (logo, title, LOPEZ, QUEZON CAMPUS) and
exactly the rows/totals the Teaching Assignment tab shows, and that a faculty
member can only ever export their own load.
"""
import csv
import io
import re

import pytest

import teaching_assignment_export as ta
from conftest import requires_db, DB_AVAILABLE

if DB_AVAILABLE:
    import app as app_module
    from database import query_db

FAC = '12079'   # a faculty account with published regular + part-time load


# ── pure helpers (no DB) ────────────────────────────────────────────────────

def test_filename_pattern_and_sanitizing():
    assert ta.export_filename('ZURBANO, Marie Andrea', '1st Semester', '2026-2027') == \
        'Teaching_Assignment_ZURBANO_Marie_Andrea_1st_Semester_2026-2027'
    assert ta.clean_filename('../../etc/pass wd.pdf') == 'etcpass_wd'
    assert ta.clean_filename('My File.docx') == 'My_File'
    assert ta.clean_filename('') == ''


def _sample():
    row = {'subjectcode': 'IT 101', 'subjectname': 'Intro', 'units': 3.0, 'year_section': 'BSIT 1',
           'time_range': '08:00 AM - 11:00 AM', 'days': 'M', 'room': 'LQ100', 'effectivity': '07/17/2026'}
    return {'source': 'official', 'fullname': 'DOE, Jane', 'employeenumber': 'E1', 'employee_type': 'Regular',
            'sem_label': '1st Semester', 'ay_label': '2026-2027', 'regular': [row], 'partTime': [],
            'totals': {'regularUnits': 3, 'partTimeUnits': 0, 'tsUnits': 0, 'totalUnits': 3,
                       'maxUnits': 21, 'availableUnits': 18}}


@pytest.mark.parametrize('fmt', ['xlsx', 'csv', 'pdf', 'docx'])
def test_every_format_builds_from_the_same_data(fmt):
    out = ta.GENERATORS[fmt](_sample(), app_module._SCH_LOGO_PATH if DB_AVAILABLE else None)
    assert isinstance(out, bytes) and len(out) > 200


def test_form_helpers():
    assert ta.compact_time_range('07:30 AM - 12:00 PM, 01:00 PM - 02:30 PM') == '7:30AM-12PM/1PM-2:30PM'
    assert ta.compact_time_range('—') == ''
    assert ta.subject_reference('DIT1') == 'T' and ta.subject_reference('BSIT 1') == 'C'
    assert ta.subject_reference('BSBIO-LQ 1-1') == 'C' and ta.subject_reference('') == ''


def test_per_day_grid_and_blank_hand_filled_parts():
    d = _sample() | {'per_day_hours': {'regular': {'Monday': 3, 'Wednesday': 1.5},
                                       'part_time': {'Saturday': 5, 'Monday': 1}}}
    rows = list(csv.reader(io.StringIO(ta.gen_csv(d).decode('utf-8-sig'))))
    grid = {r[0]: r[1:] for r in rows if r and r[0] in ('REGULAR', 'PART-TIME', 'TOTAL')}
    assert grid['REGULAR'] == ['3', '', '1.5', '', '', '', '']
    assert grid['PART-TIME'] == ['1', '', '', '', '', '5', '']
    assert grid['TOTAL'] == ['4', '', '1.5', '', '', '5', '']
    assert rows[[r[:1] for r in rows].index(['OFFICIAL TIME'])][1:] == [''] * 7
    load_row = next(r for r in rows if r and r[0] == 'IT 101')
    assert load_row[4] == 'C' and load_row[5] == '8AM-11AM' and load_row[6] == ''   # Subj. Ref. / Time / Time Code


def test_paper_form_ruled_rows_in_docx():
    from docx import Document
    doc = Document(io.BytesIO(ta.gen_docx(_sample(), None)))
    load_tables = [t for t in doc.tables if len(t.columns) == len(ta.LOAD_HEADERS)]
    assert len(load_tables) == 2 and all(len(t.rows) == 1 + ta.MIN_LOAD_ROWS for t in load_tables)


def test_empty_part_time_section_says_so():
    text = ta.gen_csv(_sample()).decode('utf-8-sig')
    assert 'PART-TIME' in text and 'Total PART-TIME,0' in text


# ── through the real route (DB) ─────────────────────────────────────────────

def _client(username=FAC, role='Faculty'):
    c = app_module.app.test_client()
    with c.session_transaction() as s:
        s.update(loggedin=True, username=username, role=role,
                 must_change_password=False, account_setup_complete=True)
    return c


@pytest.fixture
def fac():
    if not query_db("SELECT 1 FROM accounts WHERE username = %s AND isactive", [FAC]):
        pytest.skip('faculty account 12079 not present')
    c = _client()
    load = c.get('/api/faculty/my_teaching_load?source=official').get_json()
    assert load['success']
    return c, load


def _export(c, fmt, **extra):
    return c.post('/faculty/teaching-assignment/export', json={'format': fmt, 'source': 'official', **extra})


def _text_of(fmt, data):
    if fmt == 'csv':
        return data.decode('utf-8-sig')
    if fmt == 'xlsx':
        from openpyxl import load_workbook
        ws = load_workbook(io.BytesIO(data)).active
        return '\n'.join(' '.join(str(c.value) for c in r if c.value is not None) for r in ws.iter_rows())
    if fmt == 'docx':
        from docx import Document
        d = Document(io.BytesIO(data))
        return '\n'.join([p.text for p in d.paragraphs] +
                         [c.text for t in d.tables for r in t.rows for c in r.cells])
    import pdfplumber
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        return '\n'.join(p.extract_text() or '' for p in pdf.pages)


requires = pytest.mark.skipif(not DB_AVAILABLE, reason='local ASDBv11 Postgres database is not reachable')


@requires
def test_page_api_exposes_labels_and_suggested_filename(fac):
    _, load = fac
    assert load['sem_label'] in ('1st Semester', '2nd Semester', 'Summer')
    assert re.fullmatch(r'\d{4}-\d{4}', load['ay_label'])
    assert load['export_filename'].startswith('Teaching_Assignment_')
    assert load['export_filename'].endswith(f"_{load['sem_label'].replace(' ', '_')}_{load['ay_label']}")


@requires
@pytest.mark.parametrize('fmt,mime', [
    ('xlsx', 'spreadsheetml'), ('csv', 'text/csv'), ('pdf', 'application/pdf'), ('docx', 'wordprocessingml')])
def test_each_format_downloads_with_the_letterhead_and_the_same_rows(fac, fmt, mime):
    c, load = fac
    r = _export(c, fmt)
    assert r.status_code == 200, r.get_json()
    assert mime in r.mimetype
    assert f'filename="{load["export_filename"]}.{fmt}"' in r.headers['Content-Disposition']

    text = _text_of(fmt, r.data)
    for must in ('FACULTY ASSIGNMENT', 'LOPEZ, QUEZON CAMPUS', load['ay_label'],
                 'EMP NO', 'EMP STATUS', 'DEPARTMENT', 'REGULAR LOAD', 'PART-TIME', 'SUBJ. REF.',
                 'TIME CODE', 'TEACHING LOAD PER DAY (HOURS)', 'OFFICIAL TIME / ADVISING TIME',
                 'SUBJECT REFERENCE LEGEND'):
        if fmt == 'pdf':   # PDF text is read row by row, so a wrapped header's words come out apart
            assert all(w in text for w in must.split()), (fmt, must)
        else:
            assert must in ' '.join(text.split()), (fmt, must)
    for row in load['regular'] + load['partTime']:
        assert row['subjectcode'] in text and row['year_section'] in text, (fmt, row)
    assert text.count(load['regular'][0]['subjectcode']) >= 1


@requires
def test_office_formats_embed_the_pup_logo(fac):
    c, _ = fac
    from openpyxl import load_workbook
    from docx import Document
    assert len(load_workbook(io.BytesIO(_export(c, 'xlsx').data)).active._images) == 1
    assert len(Document(io.BytesIO(_export(c, 'docx').data)).inline_shapes) == 1
    assert b'/Subtype /Image' in _export(c, 'pdf').data


@requires
def test_xlsx_units_are_numbers(fac):
    c, load = fac
    from openpyxl import load_workbook
    ws = load_workbook(io.BytesIO(_export(c, 'xlsx').data)).active
    header_rows = [cell.row for cell in ws['A'] if cell.value == 'SUBJECT CODE']
    first_data = ws.cell(header_rows[0] + 1, 3).value
    assert isinstance(first_data, (int, float))


@requires
def test_local_schedule_export(fac):
    c, _ = fac
    r = c.post('/faculty/teaching-assignment/export', json={'format': 'csv', 'source': 'local'})
    assert r.status_code == 200 and 'LOCAL SCHEDULE' in r.data.decode('utf-8-sig')


@requires
def test_custom_filename_is_sanitized(fac):
    c, _ = fac
    r = _export(c, 'pdf', filename='../my load<>.pdf')
    assert 'filename="my_load.pdf"' in r.headers['Content-Disposition']


@requires
def test_only_one_valid_format_is_accepted(fac):
    c, _ = fac
    assert _export(c, 'txt').status_code == 400
    assert _export(c, '').status_code == 400


@requires
@pytest.mark.parametrize('formats', [['pdf', 'xlsx'], ['pdf', 'xlsx', 'csv', 'docx'], ['csv', 'docx', 'csv']])
def test_several_formats_download_together_as_one_zip(fac, formats):
    import zipfile
    c, load = fac
    r = c.post('/faculty/teaching-assignment/export', json={'formats': formats, 'source': 'official'})
    assert r.status_code == 200 and r.mimetype == 'application/zip'
    assert f'filename="{load["export_filename"]}.zip"' in r.headers['Content-Disposition']
    with zipfile.ZipFile(io.BytesIO(r.data)) as zf:
        names = sorted(zf.namelist())
        assert names == sorted(f'{load["export_filename"]}.{f}' for f in dict.fromkeys(formats))
        for n in names:                                  # every file is the same full export
            text = _text_of(n.rsplit('.', 1)[1], zf.read(n))
            assert 'FACULTY ASSIGNMENT' in text and load['regular'][0]['subjectcode'] in text


@requires
def test_single_format_in_the_list_is_a_plain_file(fac):
    c, load = fac
    r = c.post('/faculty/teaching-assignment/export', json={'formats': ['docx'], 'source': 'official'})
    assert r.status_code == 200 and 'wordprocessingml' in r.mimetype
    assert r.headers['Content-Disposition'].endswith('.docx"')


@requires
def test_an_invalid_format_in_the_list_is_rejected(fac):
    c, _ = fac
    for bad in (['pdf', 'txt'], [], ['']):
        assert c.post('/faculty/teaching-assignment/export', json={'formats': bad}).status_code == 400


@requires
def test_access_is_limited_to_the_logged_in_faculty(fac):
    assert app_module.app.test_client().post('/faculty/teaching-assignment/export',
                                             json={'format': 'pdf'}).status_code == 403
    assert _client('admin', 'Admin').post('/faculty/teaching-assignment/export',
                                          json={'format': 'pdf'}).status_code == 403
    # an employee number in the body is ignored — it's always the session's own load
    c, load = fac
    other = query_db("""SELECT employeenumber FROM faculty WHERE employeenumber <> %s LIMIT 1""", [FAC], one=True)
    r = _export(c, 'csv', emp_num=other['employeenumber'], employeenumber=other['employeenumber'])
    assert load['employeenumber'] in r.data.decode('utf-8-sig')
